from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from normalizer import build_normalized_transactions


def load_mapping_items(path: Path, sheet_name: str | None = None) -> tuple[str, list[dict[str, Any]]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    sheet_results = data.get("sheet_results", [])
    if not sheet_results:
        raise ValueError("Mapping file contains no sheet_results.")
    if sheet_name:
        for sheet in sheet_results:
            if sheet.get("sheet_name") == sheet_name:
                return str(sheet["sheet_name"]), sheet.get("mappings", [])
        raise ValueError(f"No mapping found for sheet: {sheet_name}")
    sheet = sheet_results[0]
    return str(sheet["sheet_name"]), sheet.get("mappings", [])


def load_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' does not exist in {path.name}")
    ws = wb[sheet_name]
    row_iter = ws.iter_rows(values_only=True)
    headers = next(row_iter, None)
    if not headers:
        raise ValueError(f"No header row found in sheet '{sheet_name}'")
    names = [str(h).strip() if h is not None else "" for h in headers]
    return [
        {names[i]: row[i] for i in range(len(names)) if names[i]}
        for row in row_iter
    ]


def main() -> None:
    if len(sys.argv) not in {3, 4}:
        print("Usage: py src/build_normalized_transactions.py <file.xlsx> <mapping.json> [sheet_name]")
        raise SystemExit(2)

    excel_path = Path(sys.argv[1])
    mapping_path = Path(sys.argv[2])
    explicit_sheet = sys.argv[3] if len(sys.argv) == 4 else None
    root = Path(__file__).resolve().parents[1]

    mapped_sheet, mapping_items = load_mapping_items(mapping_path, explicit_sheet)
    rows = load_rows(excel_path, mapped_sheet)
    normalized = build_normalized_transactions(
        rows,
        mapping_items,
        source_workbook=excel_path.name,
        source_sheet=mapped_sheet,
    )

    output_path = root / "outputs" / f"{excel_path.stem}_normalized_transactions.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved: {output_path}")
    print(f"Rows: {len(normalized['records'])}")
    print(f"Missing values: {sum(normalized['quality_summary']['missing_by_field'].values())}")
    print(f"Invalid values: {sum(normalized['quality_summary']['invalid_by_field'].values())}")


if __name__ == "__main__":
    main()
